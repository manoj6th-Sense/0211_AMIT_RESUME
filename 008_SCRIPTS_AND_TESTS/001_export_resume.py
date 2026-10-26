"""Export a Markdown résumé to HTML, DOCX (python-docx) and, with --pdf, PDF (headless Chrome).

Usage, from the project root:
    python 005_SCRIPTS/001_export_resume.py <in.md> <out_dir> <basename> [--no-docx] [--pdf]

Example:
    python 005_SCRIPTS/001_export_resume.py 001_FINAL_RESUMES/002_AMIT_MALHOTRA_FDE_RESUME_2026_ATS.md \
        001_FINAL_RESUMES/007_EXPORTS 005_AMIT_MALHOTRA_FDE_RESUME_2026 --pdf

Needs the `markdown` and `python-docx` packages (not installed by this project) and,
for --pdf, Google Chrome.
"""
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

import markdown

CHROME_CANDIDATES = (
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "google-chrome",
    "chromium",
)

CSS = """
@page { size: Letter; margin: 0.55in 0.6in; }
body { font-family: Calibri, 'Helvetica Neue', Arial, sans-serif; font-size: 10.2pt; line-height: 1.28; color: #111; }
h1 { font-size: 20pt; margin: 0 0 2pt 0; letter-spacing: 0.5px; }
h2 { font-size: 11.5pt; margin: 9pt 0 3pt 0; border-bottom: 1.2px solid #333; padding-bottom: 1pt; text-transform: uppercase; letter-spacing: 0.4px; }
h3 { font-size: 10.8pt; margin: 7pt 0 2pt 0; }
p { margin: 2pt 0 4pt 0; }
ul { margin: 1pt 0 3pt 0; padding-left: 15pt; }
li { margin: 0 0 2.2pt 0; }
table { border-collapse: collapse; width: 100%; font-size: 9.2pt; margin: 4pt 0; }
th, td { border: 1px solid #999; padding: 2pt 4pt; vertical-align: top; text-align: left; }
th { background: #eee; }
code { font-family: Menlo, monospace; font-size: 9pt; }
blockquote { margin: 3pt 0 3pt 10pt; color: #333; }
"""


def clean(text):
    """Drop the hook's stamp lines and HTML comments; they are record, not résumé."""
    meta_prefixes = ("> 🕒", "> 🗓️", "> 📋", "> **Document Version:", "> **Last Modified:", "> **Status:")
    lines = [l for l in text.splitlines() if not any(l.startswith(p) for p in meta_prefixes)]
    return re.sub(r"<!--.*?-->", "", "\n".join(lines), flags=re.S).strip() + "\n"


def write_html(text, out):
    body = markdown.markdown(text, extensions=["extra", "sane_lists"])
    out.write_text(
        f"<!doctype html><html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{body}</body></html>",
        encoding="utf-8")
    print("wrote", out)


def write_docx(text, out):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.55)
        s.left_margin = s.right_margin = Inches(0.6)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.2)

    def add_runs(par, s):
        s = s.replace("`", "")
        for i, part in enumerate(re.split(r"\*\*", s)):
            if part:
                r = par.add_run(part)
                r.bold = (i % 2 == 1)

    for raw in text.splitlines():
        l = raw.rstrip()
        if not l:
            continue
        if l.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(l[2:])
            r.bold = True
            r.font.size = Pt(20)
            p.paragraph_format.space_after = Pt(1)
        elif l.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(l[3:].upper())
            r.bold = True
            r.font.size = Pt(11.5)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            b = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "333333")):
                b.set(qn(k), v)
            pbdr.append(b)
            pPr.append(pbdr)
        elif l.startswith("### "):
            p = doc.add_paragraph()
            add_runs(p, l[4:])
            for r in p.runs:
                r.font.size = Pt(10.8)
                r.bold = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
        elif l.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, l[2:])
            p.paragraph_format.space_after = Pt(1.5)
        else:
            p = doc.add_paragraph()
            add_runs(p, l)
            p.paragraph_format.space_after = Pt(3)
    doc.save(out)
    print("wrote", out)


def find_chrome():
    for c in CHROME_CANDIDATES:
        if pathlib.Path(c).exists() or shutil.which(c):
            return c
    return None


def write_pdf(html, out):
    chrome = find_chrome()
    if not chrome:
        sys.exit("no Chrome found for --pdf; the HTML is written, print it to PDF by hand")
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run([chrome, "--headless=new", "--disable-gpu", "--no-first-run",
                        f"--user-data-dir={profile}", "--no-pdf-header-footer",
                        f"--print-to-pdf={out.resolve()}", html.resolve().as_uri()],
                       check=True, capture_output=True)
    print("wrote", out)


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    if len(args) != 3:
        sys.exit(__doc__)
    src, out_dir, base = pathlib.Path(args[0]), pathlib.Path(args[1]), args[2]
    out_dir.mkdir(parents=True, exist_ok=True)
    text = clean(src.read_text(encoding="utf-8"))
    html = out_dir / f"{base}.html"
    write_html(text, html)
    if "--no-docx" not in argv:
        write_docx(text, out_dir / f"{base}.docx")
    if "--pdf" in argv:
        write_pdf(html, out_dir / f"{base}.pdf")


if __name__ == "__main__":
    main(sys.argv[1:])
